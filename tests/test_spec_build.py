"""Specification build and content checks.

The regenerated documents must stay text-equivalent to the tracked
deliverables under ``specification/`` so that builder changes surface as
reviewable test failures rather than silent document drift.
"""

from pathlib import Path

from docx import Document

from woodland_spec import build_document, build_updated_document

REPO_ROOT = Path(__file__).resolve().parents[1]
SPEC_DIR = REPO_ROOT / "specification"

TABLE_1A_PROPERTIES = [
    "C:N Carbon to nitrogen",
    "Organic matter % of total dry weight",
    "Sand % of total dry weight",
    "Silt % of total dry weight",
    "Clay % of total dry weight",
    "Total silt and clay",
    "Acidity pH",
    "Maximum particle size",
    "Nitrogen N",
    "Phosphorus P ppm",
    "Potassium K ppm",
    "EC saturated extract at 25°C",
    "SAR saturated extract",
]

KEY_DEFAULTS = [
    "300 mm",
    "600 mm",
    "75 mm",
    "30 m",
    "5.0–6.5",
    "2.0 dS/m",
    "324 ppm",
    "1,956 ppm",
]


def _all_text(document: Document) -> list[str]:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in document.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return parts


def _tracked_text(filename: str) -> list[str]:
    return _all_text(Document(str(SPEC_DIR / filename)))


def test_original_matches_tracked_deliverable() -> None:
    regenerated = _all_text(build_document())
    assert regenerated == _tracked_text("CWH_Woodland_Growing_Medium_Specification_Original.docx")


def test_updated_matches_tracked_deliverable() -> None:
    regenerated = _all_text(build_updated_document())
    assert regenerated == _tracked_text("CWH_Woodland_Growing_Medium_Specification_Updated.docx")


def test_updated_contains_table_1a_with_all_thirteen_properties() -> None:
    document = build_updated_document()
    headings = [p.text for p in document.paragraphs]
    assert any("Table 1A" in text for text in headings)
    table_texts = [
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    ]
    for prop in TABLE_1A_PROPERTIES:
        assert prop in table_texts, prop


def test_original_does_not_contain_table_1a() -> None:
    document = build_document()
    assert not any("Table 1A" in p.text for p in document.paragraphs)


def test_key_design_defaults_present() -> None:
    text = "\n".join(_all_text(build_updated_document()))
    for token in KEY_DEFAULTS:
        assert token in text, token
