"""Export pipeline checks: Markdown source converts to a well-formed DOCX."""

import sys
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from export import export  # noqa: E402

SPEC = ROOT / "specification" / "cwh-urban-woodland-growing-medium.md"


def test_export_docx(tmp_path):
    target = export(SPEC, "docx", tmp_path)
    assert target.is_file() and target.stat().st_size > 10_000

    doc = Document(target)
    styles = [p.style.name for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]

    # Heading levels survive the round trip into Word styles.
    assert styles[0] == "Title"
    assert texts[0] == "CWH Woodland Growing Medium Specification"
    assert "Heading 1" in styles and "Heading 2" in styles
    assert "Part 6 Protection completion and measurement" in texts

    # All six tables convert; Table 1A keeps its 13 property rows.
    assert len(doc.tables) == 6
    table_1a = doc.tables[3]
    assert table_1a.rows[0].cells[0].text == "Property"
    assert len(table_1a.rows) == 14  # header plus 13 requested properties


def test_export_html(tmp_path):
    target = export(SPEC, "html", tmp_path)
    html = target.read_text(encoding="utf-8")
    assert "CWH Woodland Growing Medium Specification" in html
    assert "<table>" in html


@pytest.mark.parametrize("fmt", ["docx", "html"])
def test_export_is_deterministic(tmp_path, fmt):
    first = export(SPEC, fmt, tmp_path / "a")
    second = export(SPEC, fmt, tmp_path / "b")
    if fmt == "html":
        assert first.read_bytes() == second.read_bytes()
    else:  # DOCX zip timestamps differ; compare extracted text instead
        text = lambda p: "\n".join(x.text for x in Document(p).paragraphs)  # noqa: E731
        assert text(first) == text(second)
