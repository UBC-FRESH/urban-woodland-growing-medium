"""Build the original nine-page CWH woodland growing medium specification."""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def build_document():
    """Build and return the original specification as a python-docx Document."""
    D = Document()
    sec = D.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    sec.left_margin = sec.right_margin = Inches(0.75)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        s = D.styles[name]
        s.font.name = "Arial"
        s.font.color.rgb = RGBColor(0, 0, 0)
    D.styles["Normal"].font.size = Pt(10)
    D.styles["Normal"].paragraph_format.space_after = Pt(5)
    D.styles["Normal"].paragraph_format.line_spacing = 1.02
    D.styles["Title"].font.size = Pt(22)
    D.styles["Heading 1"].font.size = Pt(14)
    D.styles["Heading 2"].font.size = Pt(11)
    header = sec.header.paragraphs[0]
    header.text = "CWH URBAN WOODLAND  |  GROWING MEDIUM"
    header.style = "Caption"
    f = sec.footer.paragraphs[0]
    f.text = (
        "Tender specification draft  •  8 October 2026                                      Page "
    )
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    f._p.append(fld)

    def p(t):
        D.add_paragraph(t)

    def h(t):
        D.add_heading(t, 2)

    def page(t):
        D.add_page_break()
        D.add_heading(t, 1)

    def table(head, rows, widths):
        t = D.add_table(rows=1, cols=len(head))
        t.autofit = False
        for c, w in zip(t.columns, widths):
            c.width = Inches(w)
        for c, txt in zip(t.rows[0].cells, head):
            c.text = txt
        for row in rows:
            for c, txt in zip(t.add_row().cells, row):
                c.text = txt
        for i, row in enumerate(t.rows):
            pr = row._tr.get_or_add_trPr()
            ns = OxmlElement("w:cantSplit")
            pr.append(ns)
            if i == 0:
                repeat = OxmlElement("w:tblHeader")
                pr.append(repeat)
            for c in row.cells:
                tcpr = c._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ["top", "left", "bottom", "right"]:
                    b = OxmlElement("w:" + edge)
                    b.set(qn("w:val"), "single")
                    b.set(qn("w:sz"), "4")
                    b.set(qn("w:color"), "D9D9D9")
                    borders.append(b)
                tcpr.append(borders)
                sh = OxmlElement("w:shd")
                sh.set(
                    qn("w:fill"),
                    "E7E7E7" if i == 0 else ("F6F6F6" if i % 2 == 0 else "FFFFFF"),
                )
                tcpr.append(sh)
                for para in c.paragraphs:
                    para.paragraph_format.space_after = Pt(5)
                    para.paragraph_format.space_before = Pt(5)
                    for r in para.runs:
                        r.font.size = Pt(9)
                        r.bold = i == 0
        return t

    D.add_heading("CWH Woodland Growing Medium Specification", 0)
    p("Section 32 91 13  |  Constructed urban woodland on compacted fill")
    p(
        "Prepared for incorporation into landscape tender documents. Covers new trees, woodland shrubs, ferns and herbaceous plants in an open soil bed on disturbed urban ground. Includes materials, laboratory acceptance, subgrade treatment, installation and verification. All dimensions are metric."
    )
    h("Design basis and issue status")
    p(
        "Use a continuous, mineral-dominant rooting profile for a fresh-to-moist, freely drained Coastal Western Hemlock woodland. Provide a moderately organic surface layer, a lower-organic mineral rooting layer and surface wood mulch. This specification does not reproduce a natural podzol or forest floor and is not a universal specification for every CWH site series. The numerical ranges below are authored project requirements, not published CWH regulatory limits."
    )
    p(
        "This is a tender-format project specification with defined default requirements. Before issue, the Landscape Architect must complete the project schedule below and coordinate drainage and excavation with the civil, geotechnical and environmental consultants. Contractor approval submissions do not replace these design decisions. Retain the defaults unless amended expressly in the issued tender."
    )
    h("Project schedule to complete before tender issue")
    table(
        ["Item", "Default or required project entry"],
        [
            (
                "Project and authority",
                "Insert project name, site address, Contract Administrator and governing municipal requirements.",
            ),
            (
                "Ecological target",
                "Fresh-to-moist woodland. Confirm local CWH subzone, plant palette and exposure; identify separate wet or dry planting zones.",
            ),
            (
                "Soil footprint and volumes",
                "Show bed boundaries and net soil volumes on drawings. Default profile: 300 mm Layer A over 600 mm Layer B, excluding mulch.",
            ),
            (
                "Tree rooting allocation",
                "Default minimum 30 m³ of usable continuous soil per large canopy tree. Show allocation plan; sum allocations must not exceed available shared volume. Confirm against mature tree size and municipal criteria.",
            ),
            (
                "Subgrade and groundwater",
                "Issue investigation results, approved excavation limits, utility constraints, groundwater level and treatment plan. Identify areas where ripping is prohibited.",
            ),
            (
                "Drainage",
                "Issue grades, design rainfall basis, outlets, overflow route and underdrain details where required. Verify outlet permissions and capacity.",
            ),
            (
                "Environmental acceptance",
                "Environmental QP to issue applicable receiving-site soil and vapour criteria, analyte list, source-testing frequency and soil handling requirements.",
            ),
            (
                "Bid quantities",
                "Provide quantities for Layer A, Layer B, mulch, excavation, disposal, ripping and designed drainage. State any provisional quantities and contract payment rules.",
            ),
        ],
        [1.8, 5.2],
    )
    page("Part 1 General requirements")
    h("1 1 Scope and coordination")
    p(
        "Supply, test, blend, deliver and install growing medium and mulch; prepare the nonstructural landscape subgrade; protect completed work; and correct nonconforming work. Coordinate with earthworks, drainage, irrigation and planting sections. Structural soil, soil cells, roof planters, contaminated-site caps and bioretention facilities require separate details and specifications."
    )
    p(
        "Contract Administrator means the Owner’s designated representative. Landscape Architect means the landscape design professional. Environmental QP means the qualified professional responsible for environmental soil acceptance. Obtain written release at each hold point. Acceptance does not relieve the Contractor of compliance or responsibility for subsequent damage."
    )
    h("1 2 Submittals")
    p(
        "Submit at least 15 working days before proposed delivery: supplier and source locations; source history; material descriptions; blend proportions by loose volume and batching controls; anticipated quantities; representative laboratory reports; compost feedstocks and compliance records; delivery and stockpile plan; placement equipment; proposed sampling locations; trial-bed procedure; and drainage coordination records. Allow 10 working days for initial review. Do not deliver unaccepted material."
    )
    p(
        "Laboratory reports shall identify sample collector, date, lot, source, preparation, analytical method, units and extraction basis; list individual results and pass/fail against this specification. Submit chemistry and texture results no more than 60 days old for the actual production lot. Compost compliance records shall cover the supplied lot and meet the specified contractual testing and applicable regulatory frequency."
    )
    h("1 3 Laboratory and sampling requirements")
    p(
        "Use an independent laboratory experienced in landscape soils. Require ISO/IEC 17025 accreditation for environmental chemistry within its accredited scope. Where physical soil or horticultural methods are outside that scope, provide the laboratory’s documented method, quality controls and relevant experience. Contractor’s independent testing agency shall collect acceptance samples; the Owner may witness or collect verification samples."
    )
    p(
        "Define a soil production lot as one blend from unchanged component sources, not exceeding 250 m³. For each layer and each lot, collect one composite of at least 10 equal increments distributed through the pile or production stream, including its interior. Combine only material of the same layer and lot. Provide at least 10 L to the laboratory, or its required larger quantity, and retain a sealed 5 L duplicate for 90 days after acceptance. Keep samples moist and cool for biological testing; follow laboratory preservation instructions for chemistry."
    )
    p(
        "Do not composite suspected contamination hotspots. Environmental sampling shall follow the project QP’s plan and governing protocols; horticultural composites do not establish environmental compliance. Give each lot a unique identifier and connect it to all load tickets and installed locations."
    )
    h("1 4 Inspection and testing costs")
    p(
        "Include source approval, routine acceptance sampling, laboratory testing, field tests, trial bed and records in the bid. Owner-directed supplementary verification is paid by the Owner when work complies. Contractor pays investigation, removal, correction and repeat testing arising from nonconforming work, subject to the General Conditions. Notify the Contract Administrator at least two working days before each hold-point inspection."
    )
    h("1 5 Environmental and source controls")
    p(
        "Use legally sourced soil. Obtain environmental QP acceptance of existing fill proposed for retention and each imported source before mixing or delivery. Comply with applicable BC soil relocation, environmental and waste requirements. A compost quality certificate or routine nutrient test does not demonstrate receiving-site environmental suitability. Do not dilute contaminated material to obtain acceptable results."
    )
    page("Part 2 Materials")
    h("2 1 Growing medium components")
    p(
        "Layer A shall be a screened mineral sandy loam with mature plant-derived compost, meeting Table 1. Layer B shall be a compatible mineral sandy loam with only enough mature compost to meet Table 1. Use similar mineral textures in both layers. Proportion components using trial blends and tests; no fixed volumetric recipe overrides acceptance criteria."
    )
    p(
        "Mineral component: clean, noncalcareous mineral soil or blended sand and loam with stable aggregates. Sand shall be clean and washed as needed to meet salinity limits. Do not use crushed concrete, concrete washout, mortar, demolition fines, saline dredgings or uncharacterized urban excavation. Existing construction soil may be reused only after environmental and horticultural acceptance, screening and successful trial blending."
    )
    p(
        "Organic component: mature, screened compost derived from segregated leaves, yard trimmings or untreated wood residuals. Declare every feedstock. Exclude biosolids, manure, mixed municipal waste, mushroom substrate and compost containing those feedstocks. These exclusions are project material choices. Do not use peat, raw sawdust, fresh chips, undecomposed bark or forest-duff harvesting as blend components. Surface mulch is specified separately."
    )
    p(
        "Do not add lime, sulphur, fertilizer, wetting agents, biochar, hydrogel or microbial inoculants without written authorization based on testing and plant requirements. Do not apply acidification treatments to mask unsuitable alkaline construction soil. Demonstrate compliance after any authorized amendment and equilibration."
    )
    h("2 2 Compost acceptance")
    p(
        "Meet CCME Guidelines for Compost Quality 2005 Category A and applicable BC OMRR Class A compost requirements, with the more stringent corresponding requirement governing. Submit process, feedstock and lot records. Category A and Class A are separate systems; do not treat either label as evidence of the other. Apply these compost criteria to the compost ingredient, not as a substitute for testing the finished soil."
    )
    p(
        "For this contract, provide lot-specific trace-element, foreign-matter and pathogen reports even where a feedstock-specific regulatory exemption exists. Require no sharp foreign matter; total anthropogenic foreign matter greater than 2 mm shall be no more than 0.5% of compost dry mass. Comply with CCME pathogen criteria appropriate to feedstock and processing, supported by treatment records and laboratory tests."
    )
    p(
        "Compost shall have at least 21 days curing and CO₂ evolution no greater than 4 mg CO₂-C per g organic matter per day, using the applicable CCME-referenced respiration method. Require a germination and root-elongation index of at least 80% of the laboratory’s control using a documented cress bioassay. Report extraction ratio, incubation conditions and calculations. This bioassay is an additional project criterion; it does not replace stability testing."
    )
    h("2 3 Surface mulch")
    p(
        "Supply clean, arborist-produced wood chips from untreated, healthy wood, predominantly 10–75 mm, with fines smaller than 5 mm limited to 20% by loose volume. No treated or painted wood, demolition material, pallets, sharp fragments, plastic, invasive plant propagules or known diseased plant material. Declare source. Chips may contain leaves and small twigs. Do not mix mulch into the mineral soil."
    )
    p(
        "Install 75 mm settled mulch generally; use 30–50 mm around small herbaceous plants as shown in planting details. Maintain a 150 mm clear radius from tree trunks and keep mulch off shrub crowns. Retain future clean leaf fall unless it smothers small plants. Coarse woody habitat features are separate items and shall not displace required rootable soil."
    )
    page("Table 1 Growing medium laboratory acceptance")
    p(
        "Percentages are dry mass unless expressly stated. Texture fractions are normalized to 100% of the organic-matter-removed mineral fraction smaller than 2 mm. Organic matter is a percentage of the entire dry medium. Every criterion applies independently; meeting texture alone is insufficient."
    )
    table(
        ["Property and method", "Layer A upper 300 mm", "Layer B lower 600 mm"],
        [
            (
                "Mineral texture by hydrometer or pipette after organic matter removal and dispersion; sand 0.05–2 mm, silt 0.002–0.05 mm, clay <0.002 mm",
                "Sand 50–70%; silt 15–35%; clay 10–20%. Fractions total 100%.",
                "Same ranges as A. Corresponding fractions shall differ from A by no more than 10 percentage points.",
            ),
            (
                "Organic matter by loss on ignition at 550°C; report drying and ignition duration",
                "5.0–8.0%",
                "1.0–3.0%",
            ),
            ("pH in water, 1:2 dry soil mass to water volume", "5.0–6.5", "5.0–6.5"),
            (
                "Electrical conductivity in saturated paste extract at 25°C",
                "≤2.0 dS/m",
                "≤2.0 dS/m",
            ),
            (
                "Sodium adsorption ratio using Na, Ca and Mg in saturated paste extract, standard charge-based calculation",
                "≤4",
                "≤4",
            ),
            (
                "Carbon to nitrogen ratio by dry combustion; organic C divided by total N",
                "10:1–25:1",
                "Report; no pass/fail ratio at low organic C",
            ),
            (
                "Gravel 2–25 mm by sieving; screen maximum particle size",
                "≤5% of total dry medium; 100% passing 25 mm",
                "Same as A",
            ),
            (
                "Installed oven-dry bulk density by intact core and measured core volume",
                "≤1.35 Mg/m³",
                "≤1.50 Mg/m³",
            ),
            (
                "Saturated hydraulic conductivity on repacked acceptance specimens under protocol below",
                "10–100 mm/h",
                "10–100 mm/h",
            ),
            (
                "Air-filled porosity at −10 kPa on repacked acceptance specimens",
                "≥10% by volume",
                "≥10% by volume",
            ),
            (
                "Plant-available water as volumetric water at −10 kPa minus that at −1500 kPa",
                "≥12% by volume",
                "≥10% by volume",
            ),
            (
                "Total N; nitrate-N; ammonium-N; extractable P, K, Ca, Mg; hot-water-extractable B",
                "Report results and agronomic interpretation. No routine nutrient enrichment.",
                "Same as A",
            ),
            (
                "Anthropogenic debris and invasive propagules",
                "No visible construction debris, glass, metal, plastic, weeds, rhizomes or roots; no sharp matter.",
                "Same as A",
            ),
            (
                "Environmental contaminants",
                "Meet project QP-issued receiving-site criteria for the individual source and finished blend.",
                "Same as A",
            ),
        ],
        [3.3, 1.85, 1.85],
    )
    p(
        "Methods and basis: laboratory shall use the same stated preparation and extraction procedures for approval and verification. Report nutrient extraction method (for example Mehlich 3); do not compare ppm from different extraction methods or substitute 1:2 EC for saturated-paste EC. Where carbonate is detected, report organic C corrected for inorganic C. Reject a blend with evidence of phytotoxicity, anaerobic odour or active heating even if individual numerical results pass."
    )
    page("Part 3 Verification protocols and site preparation")
    h("3 1 Physical laboratory protocol")
    p(
        "Test three separately repacked specimens per layer and lot. Use rigid cylinders at least 100 mm internal diameter and 100 mm specimen height, retaining the accepted fraction up to 25 mm. Pack at the maximum installed dry density in Table 1, within ±0.03 Mg/m³. Report packing water content, mass, dimensions, actual density and gravel correction. Pre-wet slowly from below for at least 24 hours and until saturated. Test at 20 ±2°C."
    )
    p(
        "Determine saturated conductivity by constant-head permeameter, using Darcy’s law and a hydraulic gradient no greater than 1. Confirm steady flow using three readings at least 15 minutes apart differing by no more than 10%. Report head, temperature, flow, duration and individual results. Correct conductivity to 20°C. Determine retention at −10 and −1500 kPa with pressure apparatus after equilibrium; report equilibrium criterion. Calculate air-filled porosity from measured total porosity minus volumetric water at −10 kPa. Determine particle density rather than assuming 2.65 Mg/m³ for organic blends. All three replicates must meet their limits. Alternative methods require a tender addendum or written equivalency acceptance supported by paired results; changed methods shall not change acceptance limits."
    )
    h("3 2 Existing ground investigation")
    p(
        "Before soil placement, verify utility locations and permitted excavation limits. Review the issued geotechnical and environmental findings. Excavate inspection pits at least one per 250 m², minimum three per distinct bed, and at changes in fill or elevation. Examine to at least 300 mm below proposed soil formation, subject to approved excavation limits. Record fill stratification, debris, restrictive horizons, seepage, perched water, contamination indicators and adjacent structures. Pit observations supplement, rather than replace, the issued site investigation."
    )
    p(
        "Stop work where unknown utilities, suspect contamination, groundwater inconsistent with the design, structural fill or remediation barriers are encountered. Isolate the area and obtain direction under the contract. Do not rip structural fill beneath foundations, retaining structures, pavement support zones, utilities or environmental caps."
    )
    h("3 3 Subgrade treatment")
    p(
        "Remove unsuitable material within the designed rooting profile. Do not cover compacted construction soil with only a thin soil layer. In approved nonstructural landscape areas, loosen the formation to 300 mm minimum below the imported profile using excavator tines or a ripper at no more than 300 mm centres in two directions. Break clods and remove construction debris. Work only when soil fractures rather than smears. Scarify the final 100 mm to remove smooth bucket surfaces."
    )
    p(
        "Where ripping is prohibited or fails to provide a permeable, rootable formation, use the issued excavation/replacement, raised-bed or drainage detail. The Contractor shall not improvise a drainage system. Keep structural load-bearing zones separate from woodland rooting zones."
    )
    h("3 4 Subgrade acceptance")
    p(
        "At one location per 100 m², minimum three per bed, test penetration resistance through the loosened 300 mm, using a calibrated cone penetrometer near field capacity. Record moisture condition, cone geometry and depth at 50 mm increments. Require resistance below 2 MPa with no continuous restrictive band. Repeat isolated stone refusals at nearby points; persistent refusal requires excavation assessment. Dry-soil readings alone are not grounds for acceptance or rejection."
    )
    p(
        "Test infiltration at one location per 250 m², minimum three per bed and at every low point. Use a double-ring test with an inner ring at least 300 mm diameter, 50 mm maintained water head, pre-wetting and observations to a steady rate (three consecutive 30-minute intervals within 10%). Require at least 10 mm/h for the un-underdrained default profile. Record method, moisture and weather. This is a screening requirement, not a stormwater design calculation. Slower areas require the designed remedial drainage treatment and written acceptance before covering."
    )
    page("Part 4 Installation")
    h("4 1 Drainage and soil interfaces")
    p(
        "Provide the designed positive drainage and overflow route before growing medium placement. Do not create isolated soil-filled pits within impermeable fill. Where underdrains are shown, verify elevations, outlet, cleanouts and protection against fines; perform the civil specification’s flow test before covering. Do not place a blanket gravel layer or geotextile directly below growing medium unless detailed by the designer with interface performance checked. Do not place drains solely because the imported medium meets a conductivity target."
    )
    p(
        "Maintain the specified rootable profile above the design seasonal high groundwater level. Review seepage and wet-season observations with the civil consultant. The default fresh-to-moist profile does not permit persistent root-zone saturation. Flood-tolerant planting areas require a separate design."
    )
    h("4 2 Trial bed and hold points")
    p(
        "Construct a trial bed at least 10 m² using the actual accepted materials, full profile, subgrade treatment, equipment and proposed placement procedure. Demonstrate depth, blending at interfaces, settled density, infiltration and moisture condition. Complete field acceptance tests before production placement. The accepted trial establishes the workmanship standard; it does not waive testing of subsequent lots."
    )
    table(
        ["Hold point", "Required written release before proceeding"],
        [
            (
                "HP1 Material sources",
                "Source history, environmental clearance, ingredient compliance, finished blend tests and batching method accepted before delivery.",
            ),
            (
                "HP2 Formation and drainage",
                "Subgrade treatment, penetration and infiltration tests, formation survey and drainage test records accepted before covering.",
            ),
            (
                "HP3 Trial bed",
                "Trial profile, density and field infiltration pass before general placement.",
            ),
            (
                "HP4 Installed profile",
                "Each completed bed passes depth, density, chemistry verification and infiltration checks before planting or mulch conceals it.",
            ),
            (
                "HP5 Completion",
                "As-built records, corrected defects and maintenance plan accepted before completion certification.",
            ),
        ],
        [1.8, 5.2],
    )
    h("4 3 Delivery and stockpiling")
    p(
        "Deliver accepted lots separately with load ticket recording source, blend, lot, date and volume. Inspect every load for consistency, odour, weeds and debris; quarantine suspect loads. Stockpile on a clean, separated surface outside rooting and drainage areas, no higher than 2 m. Keep A and B separate. Protect against runoff, contamination, erosion and prolonged saturation with breathable covers. Do not store against tree trunks or over utilities. Inspect weekly and after heavy rain; retest stored material if contamination, heating, anaerobic odour or material change occurs."
    )
    h("4 4 Placement and settlement")
    p(
        "Use an excavator working from outside the bed or from removable load-distribution mats. Prohibit truck, loader and tracked-machine travel on exposed rooting soil. Place Layer B in loose lifts no deeper than 300 mm. Lightly settle using water and hand or low-pressure methods proven by the trial. Do not roll, vibrate or specify a Proctor compaction percentage for this planting medium."
    )
    p(
        "Place Layer A to achieve 300 mm settled thickness. Lightly intermix approximately 50 mm across the A/B boundary, without reducing the required upper layer; retain compatible texture. Loosen layer surfaces before the next lift. Do not use excessive water to slurry or compact soil. Stop placement in frozen, saturated or plastic soil, or when placement produces smearing, rutting or pumping."
    )
    p(
        "Provide at least 900 mm settled rootable medium throughout the designated woodland bed, excluding mulch and the loosened underlying formation. Include settlement allowance established by the trial; delivered loose volume is not proof of installed depth. Finished soil grade shall be within ±25 mm of the drawing elevation without reducing minimum profile depth. Shape drainage slopes as drawn; no unintended depressions. Protect building clearances and paved edges."
    )
    page("Part 5 Installed quality assurance and acceptance")
    h("5 1 Field sampling and tests")
    table(
        ["Check", "Minimum frequency and acceptance"],
        [
            (
                "Profile depths and grades",
                "Survey and probe on a grid no greater than 5 m, minimum five points per bed and each tree location. Confirm A ≥300 mm and total A+B ≥900 mm after settlement; exclude mulch. Record coordinates and levels.",
            ),
            (
                "Installed density",
                "One location per 100 m², minimum three per bed; intact cores from A at 100–200 mm and B at 450–600 mm below soil surface. Additional cores at trafficked or suspect areas. All results meet Table 1; do not average failures with passing areas.",
            ),
            (
                "Field infiltration",
                "One test per 250 m², minimum three per bed, including low points. Use the double-ring protocol in 3 4 on unmulched finished soil. Steady rate ≥10 mm/h. Laboratory conductivity and field infiltration are distinct tests.",
            ),
            (
                "Delivered blend verification",
                "Independent composite for each layer per 250 m³ delivered lot, minimum one per layer used on site. Test texture, organic matter, pH and saturated-paste EC. Retain lot traceability. Full physical panel for first lot of each layer and every changed blend/source; repeat where field failures indicate material defects.",
            ),
            (
                "Drainage behaviour",
                "Observe after a controlled 25 mm application over at least 10 m² at each representative low area, or equivalent documented rainfall. No unintended surface ponding after 24 hours and no free water within the upper 600 mm after 48 hours. Document observation wells or inspection pits. Designed wet zones require separate criteria.",
            ),
            (
                "Mulch",
                "Five depth checks per bed and one per 100 m², whichever is greater. Verify settled depth, source and trunk/crown clearance.",
            ),
            (
                "Environmental verification",
                "Follow QP-issued plan, including required discrete source and finished-blend samples. Do not substitute soil chemistry composites for environmental sampling.",
            ),
        ],
        [1.6, 5.4],
    )
    h("5 2 Planting interface")
    p(
        "Coordinate tree installation with the planting section. Set the first structural roots at final settled soil level; do not bury the root flare. Support the root ball on a firm, accepted mineral pedestal so it does not settle; preserve the continuous rootable bed around it. Backfill with the specified surrounding layer, without a separately enriched planting-hole mix. Provide drainage for the entire bed, not an unconnected drain or gravel sump beneath a tree."
    )
    p(
        "Shared soil volume may be used for tree groups only where roots can access the continuous bed. Do not count mulch, drainage aggregate, root-ball volume, utility structures or uncompromised structural fill as usable rooting medium. Confirm net allocations on as-built drawings."
    )
    h("5 3 Nonconforming work")
    p(
        "Quarantine rejected supply lots. Stop covering failed work and define affected limits using additional tests. Submit a corrective method and obtain acceptance before execution. Rework compacted soil at workable moisture or remove and replace it. For chemical or texture failures, remove the affected lot or reblend off site and retest. Do not conceal failures with mulch, add fertilizer to mask defects or dilute contaminated material. Repeat all affected acceptance tests after correction. No automatic numerical tolerance beyond the stated limits is permitted."
    )
    page("Part 6 Protection completion and measurement")
    h("6 1 Establishment protection")
    p(
        "Fence the bed against vehicle traffic and storage through construction and establishment. Repair ruts, erosion and contamination promptly. Provide temporary erosion control without plastic netting where possible, and maintain mulch clearances. After any unauthorized machine access, retest density and infiltration in the affected area before acceptance."
    )
    p(
        "Provide irrigation capable of wetting the active rooting depth, especially while the new canopy is open. Monitor at least weekly during the first growing season and during drought in the second. Base watering on measured moisture in the root ball and surrounding soil, weather and species requirements; avoid continuously saturated soil. Coordinate winter shutdown and seasonal inspections with the irrigation section."
    )
    h("6 2 Establishment period")
    p(
        "For tender pricing, provide 24 months of soil and mulch maintenance from substantial performance of the landscape work, unless the General Conditions expressly set another period. Inspect monthly during growing seasons and after major rain events; include one wet-season inspection in each year. Record settlement, ponding, erosion, weeds, mulch and soil moisture. Correct soil installation defects at Contractor cost; allocate third-party damage under the contract."
    )
    p(
        "Restore settlement greater than 25 mm relative to accepted grades, any loss of minimum depth or impeded drainage. Lift and reset affected plants where necessary rather than burying their root flares. Replenish mulch to specified depth without accumulating excess layers. Test soil before any fertilizer application. Correct the cause of plant decline associated with deficient drainage or compaction; planting replacement obligations remain governed by the planting section."
    )
    h("6 3 Completion records")
    p(
        "Provide lot register; load tickets; source approvals; all lab and field results; failed-test and correction records; profile and grade survey; tree soil allocation plan; drainage as-built drawings and test records; photographs before each layer is covered; environmental QP clearance and required relocation documentation; and a maintenance log with irrigation and mulch instructions. Identify actual component proportions and suppliers so future repairs can use compatible materials."
    )
    h("6 4 Measurement and payment")
    p(
        "Coordinate the following rules with the tender schedule and General Conditions. Measure Layer A and Layer B separately in cubic metres of accepted installed settled volume from surveyed bed area and thickness. Do not pay by truck loose volume, settlement allowance, rejected supply or rework. Measure mulch in square metres at the specified settled depth. Measure ripping in square metres of accepted treated formation; excavation and approved disposal in cubic metres under earthworks provisions."
    )
    p(
        "Include routine testing, trial work, transport, spreading, interface treatment, protection and completion records in the corresponding soil rates unless the bid schedule identifies separate items. Price underdrains, imported replacement below the default profile and QP-directed environmental work using the issued quantities or provisional items. Changed conditions shall be instructed and valued under the contract before additional work; do not transfer undefined site remediation to a fixed soil supply rate."
    )
    h("6 5 Preissue design review")
    p(
        "Before issuing this section for bids, the Landscape Architect shall confirm plant suitability for pH 5.0–6.5 and the selected moisture regime, adequate connected soil volume and growing-season water supply. Civil and geotechnical consultants shall resolve drainage, settlement, structures and utility interfaces. The environmental QP shall complete the contaminant schedule. List exact editions of any additional standards incorporated into the contract and resolve conflicts expressly. Do not issue an undefined drainage outlet, unknown rootable depth or missing environmental acceptance schedule as a contractor design assumption."
    )
    page("Appendix A Technical basis and references")
    p(
        "This appendix explains the specification’s basis; it is not an additional scope of work. Technical sources were checked on 8 October 2026. The requirements in Parts 1–6 are original project clauses. No source below publishes this entire specification or a single official CWH woodland blend."
    )
    h("Ecological basis")
    p(
        "BC’s CWH description identifies a wide coastal climatic range, including drier inner-coast and wetter outer-coast conditions. The western hemlock species account describes acid organic regeneration substrates, small balanced nutrient supplies and continuous moisture, with sensitivity to warm dry conditions. These support the selected acidic, moisture-retentive woodland concept, while requiring local site and species confirmation."
    )
    p(
        "Province of British Columbia. Coastal western hemlock zone. https://www2.gov.bc.ca/gov/content/industry/forestry/managing-our-forest-resources/silviculture/stocking-standards/tree-species-select/climatic-zones-bc/cwh-zone"
    )
    p(
        "Province of British Columbia. Western hemlock species compendium. https://www2.gov.bc.ca/gov/content/industry/forestry/managing-our-forest-resources/silviculture/stocking-standards/tree-species-compendium-index/western-hemlock"
    )
    h("Compaction and physical acceptance")
    p(
        "USDA NRCS explains that bulk density depends on texture and organic matter and that construction traffic, particularly on wet soil, impairs pore space, drainage and rooting. The density caps, porosity and water-retention targets in this specification are conservative project selections for the specified sandy-loam blends; they are not quoted NRCS woodland limits. The 900 mm profile and 30 m³ tree allocation are design defaults requiring confirmation against the actual trees and local standards."
    )
    p(
        "USDA Natural Resources Conservation Service. Soil Quality Indicators Bulk Density. June 2008. https://www.nrcs.usda.gov/sites/default/files/2023-01/Soil%20Quality-Indicators-Bulk%20Density.pdf"
    )
    h("Compost and environmental compliance")
    p(
        "CCME 2005 provides Category A compost criteria and maturity/stability approaches. BC OMRR has separate Class A criteria; the contract requires compliance with both. The supplied compost must be tested before blending. BC soil relocation guidance establishes source and receiving-site obligations; the environmental QP must determine the applicable criteria and protocols for this project. The compost provisions do not authorize contaminated fill reuse."
    )
    p(
        "Canadian Council of Ministers of the Environment. Guidelines for Compost Quality. 2005, PN 1340, sections 3 and 4. https://ccme.ca/en/res/compostgdlns_1340_e.pdf"
    )
    p(
        "British Columbia. Organic Matter Recycling Regulation, B.C. Reg. 18/2002, current consolidation. https://www.bclaws.gov.bc.ca/civix/document/id/complete/statreg/18_2002"
    )
    p(
        "Province of British Columbia. Soil relocation. https://www2.gov.bc.ca/gov/content/environment/air-land-water/site-remediation/soil-relocation"
    )
    h("Municipal contract coordination")
    p(
        "Vancouver Park Development Standards provides municipal design/construction guidance and downloadable specification sections. Use the applicable authority’s issued project requirements for coordination. Municipal numerical soil limits have not been copied into this bespoke specification. The Vancouver GRI planting guidelines were not used."
    )
    p(
        "City of Vancouver. Park Development Standards. https://vancouver.ca/parks-recreation-culture/park-development-standards.aspx"
    )
    for style in D.styles:
        if style.type == 1:
            style.font.color.rgb = RGBColor(0, 0, 0)
            for el in list(style.element.iter(qn("w:pBdr"))):
                el.getparent().remove(el)
    for para in D.paragraphs:
        for el in list(para._p.iter(qn("w:pBdr"))):
            el.getparent().remove(el)
    return D
